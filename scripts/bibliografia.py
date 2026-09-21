"""Gera artigos/referencias_citadas.xlsx a partir do referencias.bib e das citações nos .tex.

Roda sozinho e não depende do pipeline: lê apenas a pasta dissertacao/. As colunas de leitura
(papel na dissertação, situação de acesso) ficam declaradas aqui, e não no .bib, porque são
anotações do trabalho e não campos bibliográficos.

    uv run python scripts/bibliografia.py

Depois de gerar, o arquivo passa por um recálculo no LibreOffice, se ele estiver instalado,
para que as fórmulas de resumo tenham valor gravado.
"""
import collections
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
DISSERTACAO = RAIZ / 'dissertacao'
SAIDA = RAIZ / 'artigos' / 'referencias_citadas.xlsx'
FONTE = 'Arial'

CMD = re.compile(r'\\(?:cite|citeonline|parencite|textcite|citeauthor|citeyear|footcite)'
                 r'\*?(?:\[[^\]]*\])*\{([^}]*)\}')

CAPITULOS = {
    '01-introducao': 'Cap. 1 Introdução',
    '02-referencial-teorico': 'Cap. 2 Referencial',
    '03-procedimentos-metodologicos': 'Cap. 3 Método',
    '04-analise-e-discussao': 'Cap. 4 Resultados',
    '05-consideracoes-finais': 'Cap. 5 Agenda',
}

TIPOS = {'article': 'Artigo em periódico', 'report': 'Relatório / working paper',
         'inproceedings': 'Trabalho em anais', 'online': 'Recurso eletrônico',
         'book': 'Livro', 'incollection': 'Capítulo de livro', 'thesis': 'Tese'}

PAPEL = {
    'felten2021': 'Fonte do AIOE, o tratamento principal do trabalho.',
    'eloundou2024': 'Medida de exposição por tarefas; substituição prevista como robustez.',
    'brynjolfsson2025': 'Adoção observada em uma firma; contraste de desenho e de magnitude.',
    'humlum2025': 'Efeitos nulos e precisos com registros administrativos; calibra a expectativa.',
    'acemoglu2019': 'Quadro de tarefas: efeitos de deslocamento e de recomposição.',
    'autor2025': 'Dimensão de especialização das tarefas remanescentes; sinais opostos entre '
                 'rendimento e emprego.',
    'acemoglu2020': 'Robôs e exposição diferencial; origem da estratégia de identificação.',
    'dingel2020': 'Índice de teletrabalhabilidade, usado como controle interagido com o período.',
    'callaway2024': 'Diferenças em diferenças com tratamento contínuo; define o parâmetro.',
    'rambachan2023': 'Sensibilidade a desvios das tendências paralelas.',
    'ulyssea2020': 'Informalidade como fenômeno heterogêneo; justifica o desfecho de formalidade.',
    'meghir2015': 'Busca e emparelhamento; fronteira formal-informal atravessada com frequência.',
    'higano2023': 'Mobilidade ocupacional e rendimentos na PME; contrafactual de rotatividade.',
    'wroblevski2024': 'Matrizes de transição com a PNAD Contínua; mobilidade na informalidade.',
    'osorio2022': 'Identificação longitudinal dos painéis da PNAD Contínua.',
    'ibge2022tratamento': 'Mudança no modo de coleta em 2020 e 2021; base da ressalva de medida.',
    'basedosdados_pnadc': 'Fonte dos microdados da PNAD Contínua usados na estimação.',
    'massenkoff2026': 'Exposição observada e DiD sobre desemprego; inclusão em avaliação.',
    'ahn2026': 'Fluxos do mercado de trabalho por exposição e adoção; margem de ajuste no '
               'lado da contratação. Inclusão em avaliação.',
    'yin2026': 'Instabilidade das medidas de exposição geradas por modelos de linguagem; '
               'sustenta a ressalva sobre a medida de tratamento. Inclusão em avaliação.',
}

ACESSO = {
    'acemoglu2019': 'Aberto (JEP)',
    'higano2023': 'Aberto (SciELO)',
    'osorio2022': 'Aberto (repositório do Ipea)',
    'wroblevski2024': 'Aberto (anais da ANPEC)',
    'ibge2022tratamento': 'Aberto (biblioteca do IBGE)',
    'basedosdados_pnadc': 'Aberto (site, sem PDF)',
    'massenkoff2026': 'Aberto (site da Anthropic)',
    'callaway2024': 'Restrito no NBER; há versão aberta em pré-print',
    'felten2021': 'Aberto no SSRN (versão de trabalho); publicado restrito no SMJ',
    'ahn2026': 'Aberto (página do autor e site de conferência do NBER)',
    'yin2026': 'Aberto (NBER); também em pré-print no SSRN',
    'autor2025': 'Restrito na JEEA; versão aberta no NBER e na página do MIT',
    'eloundou2024': 'Restrito na Science; há versão aberta em pré-print',
    'humlum2025': 'Restrito no NBER; há versão aberta na página dos autores',
    'brynjolfsson2025': 'Restrito no QJE; há versão aberta em working paper',
}

# Obras ainda não citadas no texto que devem constar da planilha.
EXTRA = {
    'massenkoff2026': {
        'tipo': 'report', 'author': 'Massenkoff, Maxim and McCrory, Peter',
        'title': 'Labor market impacts of AI',
        'subtitle': 'a new measure and early evidence',
        'institution': 'Anthropic', 'year': '2026',
        'url': 'https://www.anthropic.com/research/labor-market-impacts',
    },
    'yin2026': {
        'tipo': 'report', 'author': 'Yin, Michelle and Vu, Hoa and Persico, Claudia',
        'title': 'How (un)stable are LLM occupational exposure scores?',
        'subtitle': 'evidence from multi-model replication',
        'institution': 'National Bureau of Economic Research', 'year': '2026',
        'number': '35110', 'url': 'https://www.nber.org/papers/w35110',
    },
    'ahn2026': {
        'tipo': 'report', 'author': 'Ahn, Hie Joo and Carollo, Nicholas A.',
        'title': 'Artificial intelligence and labor market reallocation',
        'institution': 'Board of Governors of the Federal Reserve System', 'year': '2026',
        'note': 'Versão preliminar de 16 de setembro de 2026',
        'url': 'https://www.ncarollo.net/',
    },
}

ARQUIVOS = {'massenkoff2026': 'massenkoff_mccrory_2026_labor_market_impacts_ai.pdf'}

# Qual versão do texto está guardada na pasta. Vários periódicos são restritos e o que se
# consegue é a versão de trabalho, que pode diferir da publicada em números e redação.
VERSAO = {
    'acemoglu2019': 'NBER WP 25684; publicado no JEP',
    'acemoglu2020': 'NBER WP 23285; publicado no JPE',
    'brynjolfsson2025': 'NBER WP 31161; publicado no QJE',
    'callaway2024': 'Pré-print arXiv 2107.02637v8, dez. 2025',
    'dingel2020': 'NBER WP 26948; publicado no JPubE',
    'eloundou2024': 'Pré-print arXiv 2303.10130v5, 2023; publicado na Science em 2024',
    'higano2023': 'Versão publicada (RBE)',
    'humlum2025': 'NBER WP 33777',
    'ibge2022tratamento': 'Versão publicada (IBGE)',
    'massenkoff2026': 'Versão publicada (Anthropic)',
    'meghir2015': 'NBER WP 18347; publicado na AER',
    'osorio2022': 'Versão publicada (Ipea)',
    'rambachan2023': 'Versão publicada (ReStud)',
    'ulyssea2020': 'Versão aceita (repositório da UCL)',
    'wroblevski2024': 'Versão publicada (anais da ANPEC)',
    'felten2021': 'Versão de trabalho do SSRN; publicado no SMJ',
    'ahn2026': 'Versão preliminar de 16 set. 2026',
    'yin2026': 'NBER WP 35110',
    'autor2025': 'NBER WP 33941; publicado na JEEA',
}


def citacoes():
    uso = collections.defaultdict(list)
    for tex in sorted(DISSERTACAO.rglob('*.tex')):
        if 'tabelas' in tex.parts:
            continue
        for i, linha in enumerate(tex.read_text(encoding='utf-8').splitlines(), 1):
            if linha.lstrip().startswith('%'):
                continue
            for m in CMD.finditer(linha):
                for chave in m.group(1).split(','):
                    uso[chave.strip()].append(tex.stem)
    return uso


def entradas_bib():
    texto = (DISSERTACAO / 'referencias.bib').read_text(encoding='utf-8')
    entradas = {}
    for m in re.finditer(r'@(\w+)\{([^,]+),(.*?)\n\}', texto, re.S):
        corpo = m.group(3) + '\n'
        campos = {c.group(1).lower(): ' '.join(c.group(2).split())
                  for c in re.finditer(r'(\w+)\s*=\s*\{(.*?)\}\s*,?\s*\n(?=\s*\w+\s*=|\s*$)',
                                       corpo, re.S)}
        entradas[m.group(2).strip()] = dict(tipo=m.group(1), **campos)
    return entradas


def linha(chave, entrada, onde):
    autores = entrada.get('author', '').replace(' and ', '; ').replace('{', '').replace('}', '')
    veiculo = (entrada.get('journaltitle') or entrada.get('journal') or entrada.get('institution')
               or entrada.get('eventtitle') or entrada.get('publisher') or '')
    detalhe = [p for p in (f"v. {entrada['volume']}" if entrada.get('volume') else '',
                           f"n. {entrada['number']}" if entrada.get('number') else '',
                           f"p. {entrada['pages'].replace('--', '-')}" if entrada.get('pages') else '')
               if p]
    capitulos = []
    for stem in onde:
        rotulo = CAPITULOS.get(stem, stem)
        if rotulo not in capitulos:
            capitulos.append(rotulo)
    arquivo = ARQUIVOS.get(chave, f'{chave}.pdf')
    if not (RAIZ / 'artigos' / arquivo).exists():
        arquivo = ''
    return {
        'Chave': chave,
        'Autores': autores,
        'Ano': int(entrada['year']) if entrada.get('year', '').isdigit() else entrada.get('year', ''),
        'Título': entrada.get('title', ''),
        'Subtítulo': entrada.get('subtitle', ''),
        'Tipo': TIPOS.get(entrada.get('tipo', ''), entrada.get('tipo', '')),
        'Veículo': veiculo,
        'Detalhes': ', '.join(detalhe),
        'DOI': entrada.get('doi', ''),
        'Link': entrada.get('url', '') or (f"https://doi.org/{entrada['doi']}"
                                           if entrada.get('doi') else ''),
        'Citações': len(onde),
        'Onde aparece': '; '.join(capitulos) if capitulos else 'Ainda não citado',
        'Papel na dissertação': PAPEL.get(chave, ''),
        'PDF na pasta': arquivo or ('não se aplica' if entrada.get('tipo') == 'online'
                                    else 'falta baixar'),
        'Versão do arquivo': VERSAO.get(chave, ''),
        'Acesso': ACESSO.get(chave, 'Restrito (periódico)'),
    }


LARGURAS = {'Chave': 20, 'Autores': 38, 'Ano': 6, 'Título': 42, 'Subtítulo': 34, 'Tipo': 22,
            'Veículo': 30, 'Detalhes': 20, 'DOI': 28, 'Link': 42, 'Citações': 9,
            'Onde aparece': 34, 'Papel na dissertação': 52, 'PDF na pasta': 24,
            'Versão do arquivo': 38, 'Acesso': 34}


def escreve(linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Referências citadas'
    cabecalho = list(linhas[0])
    ws.append(cabecalho)
    for item in linhas:
        ws.append([item[c] for c in cabecalho])
    ultima = ws.max_row

    resumo = ultima + 2
    for deslocamento, (rotulo, formula) in enumerate([
            ('Total de obras', f'=COUNTA(A2:A{ultima})'),
            ('Total de citações no texto', f'=SUM(K2:K{ultima})'),
            ('PDFs já na pasta', f'=COUNTIF(N2:N{ultima},"*.pdf")')]):
        ws.cell(resumo + deslocamento, 1, rotulo).font = Font(name=FONTE, bold=True)
        ws.cell(resumo + deslocamento, 2, formula).font = Font(name=FONTE)
    ws.cell(resumo + 5, 1,
            'Gerado por scripts/bibliografia.py a partir de dissertacao/referencias.bib e das '
            'chaves de citação nos arquivos .tex. As colunas Papel na dissertação e Acesso são '
            'anotações mantidas no próprio script.').font = Font(name=FONTE, italic=True, size=9)

    azul = PatternFill('solid', fgColor='1F3864')
    for col, nome in enumerate(cabecalho, 1):
        c = ws.cell(1, col)
        c.font = Font(name=FONTE, bold=True, color='FFFFFF')
        c.fill = azul
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = LARGURAS.get(nome, 18)
    for fila in ws.iter_rows(min_row=2, max_row=ultima, max_col=len(cabecalho)):
        for c in fila:
            c.font = Font(name=FONTE, size=10)
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cabecalho))}{ultima}'
    ws.row_dimensions[1].height = 28

    SAIDA.parent.mkdir(exist_ok=True)
    wb.save(SAIDA)


def recalcula():
    """Reabre o arquivo no LibreOffice para gravar o valor das fórmulas de resumo."""
    soffice = Path('C:/Program Files/LibreOffice/program/soffice.exe')
    if not soffice.exists():
        soffice = shutil.which('soffice')
    if not soffice:
        print('LibreOffice não encontrado: fórmulas ficam sem valor gravado.')
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([str(soffice), '--headless', '--norestore', '--calc', '--convert-to',
                        'xlsx:Calc MS Excel 2007 XML', '--outdir', tmp, str(SAIDA)],
                       check=True, capture_output=True, timeout=180)
        convertido = Path(tmp) / SAIDA.name
        valores = load_workbook(convertido, data_only=True).active
        erros = [c.coordinate for fila in valores.iter_rows() for c in fila
                 if isinstance(c.value, str) and c.value.startswith('#')]
        if erros:
            sys.exit(f'Fórmulas com erro em {erros}')
        shutil.copy2(convertido, SAIDA)


if __name__ == '__main__':
    uso = citacoes()
    entradas = entradas_bib()
    chaves = sorted(uso, key=lambda k: (entradas[k].get('author', ''), entradas[k].get('year', '')))
    dados = [linha(k, entradas[k], uso[k]) for k in chaves]
    dados += [linha(k, e, uso.get(k, [])) for k, e in EXTRA.items() if k not in uso]
    escreve(dados)
    recalcula()
    print(f'{SAIDA.relative_to(RAIZ)}: {len(dados)} obras, '
          f'{sum(d["Citações"] for d in dados)} citações no texto')
